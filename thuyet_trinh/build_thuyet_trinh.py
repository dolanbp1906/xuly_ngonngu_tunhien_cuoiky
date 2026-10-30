#!/usr/bin/env python3
"""Sinh PPTX báo cáo và DOCX kịch bản thuyết trình ViKAG-Faith.

Số liệu lấy từ BAI_BAO.md, eval/summary.txt, eval/viedureg_systems.txt,
eval/public_report.txt (chạy offline, không gọi mô hình sinh).
"""

from __future__ import annotations

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRGB
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn as pqn
from pptx.util import Emu, Inches, Pt

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn as dqn
from docx.shared import Cm, Pt as DPt, RGBColor as DocRGB

OUT_DIR = __file__.rsplit("/", 1)[0]
PPTX_PATH = f"{OUT_DIR}/ViKAG-Faith_bao_cao.pptx"
DOCX_PATH = f"{OUT_DIR}/Kich_ban_thuyet_trinh.docx"

BG = "FFF8F3"
INK = "1C2430"
MUTED = "5C6B7A"
TEAL = "0E9A86"
TEAL_DEEP = "0B6E62"
CORAL = "FF6B57"
CORAL_DEEP = "E24B3C"
AMBER = "F0A202"
AMBER_INK = "A16207"
SKY = "3D7EFF"
GRAPE = "6E63F5"
WHITE = "FFFFFF"
MINT = "E7F7F3"
PEACH = "FFECEA"
SUN = "FFF6E4"
SKY_BG = "EAF1FF"
GRAPE_BG = "F1EFFF"
LINE = "F0E4DA"
TRACK = "F3EBE4"

FONT = "Arial"
W = 13.333333
H = 7.5
TOTAL = 15


def rgb(h: str) -> PptRGB:
    return PptRGB.from_string(h)


def set_run(run, size, color, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = False
    run.font.color.rgb = rgb(color)
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        node = rpr.find(pqn(tag))
        if node is None:
            node = etree.SubElement(rpr, pqn(tag))
        node.set("typeface", FONT)


def no_line(shape):
    shape.line.fill.background()


def rect(slide, x, y, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(fill)
    no_line(sh)
    return sh


def rrect(slide, x, y, w, h, fill, radius=0.12):
    sh = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(fill)
    no_line(sh)
    try:
        sh.adjustments[0] = radius
    except Exception:
        pass
    return sh


def oval(slide, x, y, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(fill)
    no_line(sh)
    return sh


def _align(name):
    return {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[name]


def _anchor(name):
    return {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[name]


def fill_tf(shape, blocks, anchor="top", ml=0.14, mr=0.12, mt=0.08, mb=0.06):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = None
    tf.anchor = _anchor(anchor)
    tf.margin_left = Inches(ml)
    tf.margin_right = Inches(mr)
    tf.margin_top = Inches(mt)
    tf.margin_bottom = Inches(mb)
    for i, block in enumerate(blocks):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = _align(block.get("align", "left"))
        p.space_before = Pt(block.get("before", 0))
        p.space_after = Pt(block.get("after", 2))
        p.line_spacing = block.get("leading", 1.0)
        parts = block["parts"]
        if isinstance(parts, str):
            parts = [(parts, block.get("bold", False), block.get("color", INK), block.get("size", 15))]
        for text, bold, color, size in parts:
            run = p.add_run()
            run.text = text
            set_run(run, size, color, bold)
    return tf


def tb(slide, x, y, w, h, blocks, anchor="top", ml=0.02, mr=0.02, mt=0.0, mb=0.0):
    sh = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    fill_tf(sh, blocks, anchor=anchor, ml=ml, mr=mr, mt=mt, mb=mb)
    return sh


def P(text, size=15, color=INK, bold=False, align="left", before=0, after=2, leading=1.0):
    return {
        "parts": text if isinstance(text, list) else [(text, bold, color, size)],
        "align": align,
        "before": before,
        "after": after,
        "leading": leading,
    }


def R(text, bold=False, color=INK, size=15):
    return (text, bold, color, size)


def blank(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, W, H, BG)
    return slide


def footer(slide, n):
    rect(slide, 0.48, 7.12, 12.38, 0.015, LINE)
    tb(
        slide, 0.48, 7.16, 8.5, 0.28,
        [P("ViKAG-Faith   ·   Xử lý ngôn ngữ tự nhiên", 11, MUTED)],
        anchor="middle",
    )
    pill = rrect(slide, 12.15, 7.16, 0.7, 0.28, TEAL, 0.4)
    fill_tf(
        pill,
        [P(f"{n:02d}", 11, WHITE, True, "center")],
        anchor="middle", ml=0.02, mr=0.02, mt=0.0, mb=0.0,
    )


def header(slide, kicker, title):
    rect(slide, 0, 0, 10.4, 0.08, TEAL)
    rect(slide, 10.4, 0, W - 10.4, 0.08, CORAL)
    tb(slide, 0.5, 0.18, 12.3, 0.26, [P(kicker.upper(), 12, TEAL, True)])
    tb(slide, 0.48, 0.44, 12.35, 0.52, [P(title, 26, INK, True)])


def notes(slide, sec):
    body = sec["say"].strip()
    extra = sec.get("point", "")
    bridge = sec.get("bridge", "")
    text = body
    if extra:
        text += "\n\nChỉ vào: " + extra
    if bridge:
        text += "\n\nChuyển: " + bridge
    slide.notes_slide.notes_text_frame.text = text


def bar(slide, x, y, w, h, value, fill, track=TRACK):
    rrect(slide, x, y, w, h, track, 0.5)
    fw = max(0.08, w * max(0.0, min(1.0, value)))
    rrect(slide, x, y, fw, h, fill, 0.5)


# ---------------------------------------------------------------------------
# Slides
# ---------------------------------------------------------------------------

def slide_01(prs):
    s = blank(prs)
    rrect(s, 0.36, 0.32, 6.35, 6.86, TEAL, 0.08)
    oval(s, 4.85, 5.15, 1.55, 1.55, "0C7C70")
    oval(s, 0.62, 5.95, 0.7, 0.7, "13B39A")
    tb(s, 0.72, 0.62, 5.5, 0.32, [P("PROJECT 4   ·   CUỐI KỲ", 13, "D7FFF6", True)])
    tb(s, 0.7, 1.15, 5.6, 0.85, [P("ViKAG-Faith", 44, WHITE, True)])
    tb(
        s, 0.72, 2.15, 5.35, 1.35,
        [P("Đánh giá tính trung thực mức mệnh đề cho hỏi đáp quy chế đào tạo tiếng Việt.", 18, "F3FFFC", leading=1.05)],
    )
    quote = rrect(s, 0.72, 3.7, 5.15, 1.15, "0C7C70", 0.14)
    fill_tf(
        quote,
        [P("Không chấm cả câu.\nChấm từng mệnh đề.", 20, WHITE, True, leading=1.05)],
        anchor="middle", ml=0.22, mr=0.16, mt=0.08, mb=0.06,
    )
    tb(s, 0.72, 5.15, 4.0, 0.7, [P("Xử lý ngôn ngữ tự nhiên\n30 tháng 10 năm 2026", 14, "E7FFF9", leading=1.05)])

    tb(s, 7.05, 0.48, 5.7, 0.3, [P("BUỔI NÓI ĐI QUA", 12, TEAL, True)])
    items = [
        ("01", TEAL, MINT, "Vì sao một câu xuôi vẫn sai", "Đổi số, đổi hệ quả, bịa thêm."),
        ("02", SKY, SKY_BG, "Cách chấm trên đồ thị điều–khoản", "Ba nhãn và thước GGF."),
        ("03", GRAPE, GRAPE_BG, "Số đã chạy trên CPU", "Truy hồi, verifier, vòng sửa."),
        ("04", CORAL, PEACH, "Demo ba câu quy chế", "Đúng, nói ngược, rồi sửa."),
    ]
    y = 0.95
    for num, accent, bg, title, sub in items:
        card = rrect(s, 7.0, y, 5.9, 1.22, WHITE, 0.12)
        fill_tf(card, [P("")], ml=0, mr=0, mt=0, mb=0)
        oval(s, 7.22, y + 0.32, 0.58, 0.58, accent)
        tb(s, 7.22, y + 0.42, 0.58, 0.4, [P(num, 12, WHITE, True, "center")], anchor="middle")
        tb(s, 8.0, y + 0.22, 4.65, 0.42, [P(title, 16, INK, True)])
        tb(s, 8.0, y + 0.64, 4.65, 0.38, [P(sub, 13, MUTED)])
        y += 1.36
    tb(
        s, 7.05, 6.5, 5.8, 0.45,
        [P("Khoảng 12 phút nói  ·  rồi demo trực tiếp", 13, MUTED)],
    )
    return {
        "n": 1,
        "kicker": "Mở đầu",
        "title": "ViKAG-Faith",
        "seconds": 35,
        "cut": "ngắn",
        "goal": "Chào và khóa một câu: chấm từng mệnh đề, không chấm cả câu.",
        "say": (
            "Em chào thầy cô. Em trình bày Project 4, tên ngắn là ViKAG-Faith. "
            "Đề tài đánh giá tính trung thực của câu trả lời trên quy chế đào tạo tiếng Việt. "
            "Điểm khác là em không chấm cả câu bằng một số. Em chấm từng mệnh đề: "
            "mệnh đề nào được khoản kéo theo, mệnh đề nào nói ngược khoản, "
            "và mệnh đề nào thêm thông tin không có trong nguồn. "
            "Phần sau em đi từ ví dụ, cách làm, rồi các số đã chạy trên máy CPU."
        ),
        "point": "Câu trên tấm nền xanh: Không chấm cả câu. Chấm từng mệnh đề.",
        "bridge": "Em bắt đầu bằng một ví dụ sinh viên gặp thật.",
    }


def slide_02(prs):
    s = blank(prs)
    header(s, "Vấn đề", "Câu nghe xuôi vẫn có thể sai quy chế")
    cards = [
        (CORAL_DEEP, PEACH, "01", "Đổi số", "Nguồn ghi vượt quá 24 tín chỉ.", "Câu trả lời nói 30."),
        (CORAL_DEEP, "FFF1EE", "02", "Đổi hệ quả", "Số 24 vẫn đúng, nhưng cảnh báo học tập", "bị nói thành buộc thôi học."),
        (AMBER_INK, SUN, "03", "Bịa thêm", "Khoản không nói tiền thưởng.", "Câu thêm 3 triệu mỗi học kỳ."),
    ]
    x = 0.48
    for accent, bg, num, title, a, b in cards:
        rrect(s, x, 1.25, 4.0, 3.55, bg, 0.1)
        rect(s, x, 1.25, 4.0, 0.1, accent)
        tb(s, x + 0.22, 1.5, 3.5, 0.35, [P(num, 13, accent, True)])
        tb(s, x + 0.22, 1.9, 3.55, 0.5, [P(title, 24, INK, True)])
        tb(s, x + 0.22, 2.6, 3.55, 0.85, [P(a, 15, MUTED, leading=1.05)])
        tb(s, x + 0.22, 3.5, 3.55, 0.9, [P(b, 16, INK, True, leading=1.05)])
        x += 4.16
    banner = rrect(s, 0.48, 5.05, 12.38, 1.85, WHITE, 0.1)
    fill_tf(
        banner,
        [
            P("Vì sao thước cũ không đủ", 16, TEAL_DEEP, True, after=6),
            P(
                "F1 của cả câu, và điểm một giám khảo chấm cả câu, gộp ba lỗi này thành một số. "
                "Người đọc không thấy mệnh đề nào được khoản kéo theo, mệnh đề nào nói ngược, mệnh đề nào bịa.",
                16, INK, leading=1.05, after=4,
            ),
            P("Lỗi đổi hệ quả nguy hiểm hơn lỗi thiếu chi tiết, nên contradicted đứng riêng, không gộp với unsupported.", 14, MUTED, leading=1.02),
        ],
        ml=0.28, mr=0.24, mt=0.16, mb=0.1,
    )
    footer(s, 2)
    return {
        "n": 2,
        "kicker": "Vấn đề",
        "title": "Câu nghe xuôi vẫn có thể sai quy chế",
        "seconds": 60,
        "cut": "ngắn",
        "goal": "Ba kiểu sai phải tách được. Nhấn kiểu đổi hệ quả.",
        "say": (
            "Quy chế là văn bản sinh viên dùng khi quyết định việc cụ thể, ví dụ nợ bao nhiêu tín chỉ thì bị cảnh báo. "
            "Một câu nghe xuôi vẫn sai theo ba kiểu. "
            "Một, đổi số: khoản nói vượt quá 24, câu nói 30. "
            "Hai, đổi hệ quả: số 24 vẫn đúng, nhưng cảnh báo học tập bị nói thành buộc thôi học. "
            "Lỗi này nguy hiểm hơn, vì người đọc tưởng mức xử lý nặng hơn quy chế. "
            "Ba, bịa thêm: khoản không nói tiền, câu thêm thưởng ba triệu mỗi học kỳ. "
            "Điểm F1 của cả câu không tách được ba kiểu này. "
            "Đề tài chấm từng mệnh đề để thấy khoản nào kéo theo, mệnh đề nào nói ngược, mệnh đề nào thêm ngoài nguồn."
        ),
        "point": "Thẻ giữa, dòng buộc thôi học.",
        "bridge": "Từ ví dụ đó, em đặt ba câu hỏi, và nói hướng trả lời trước.",
    }


def slide_03(prs):
    s = blank(prs)
    header(s, "Câu hỏi nghiên cứu", "Ba câu hỏi, trả lời ngắn trước")
    rows = [
        ("RQ1", TEAL, MINT, "Thước tự động khớp người tới đâu?",
         "Thước có xét số và hệ quả bám nhãn người hơn thước đếm từ. Tập người còn quá nhỏ để kết luận tương quan ổn định."),
        ("RQ2", SKY, SKY_BG, "Từng khoản và rule số có hơn gộp cả điều?",
         "Có, trên 14 mệnh đề đã biết: macro-F1 1,000, bỏ rule số còn 0,436. Trên toàn văn, lớp contradicted chỉ đúng khoảng 0,15."),
        ("RQ3", CORAL, PEACH, "Vòng sửa có giảm câu nói ngược?",
         "Có, khi thay mảnh lệch bằng đúng khoản: contradicted về 0. Chưa phải mô hình sinh viết lại, chưa đo F1 câu viết."),
    ]
    y = 1.22
    for code, accent, bg, q, a in rows:
        rrect(s, 0.48, y, 12.38, 1.72, WHITE, 0.1)
        badge = rrect(s, 0.7, y + 0.5, 1.15, 0.7, accent, 0.18)
        fill_tf(badge, [P(code, 16, WHITE, True, "center")], anchor="middle", ml=0.02, mr=0.02, mt=0, mb=0)
        tb(s, 2.1, y + 0.18, 10.4, 0.48, [P(q, 20, INK, True)])
        tb(s, 2.1, y + 0.72, 10.4, 0.8, [P(a, 15, MUTED, leading=1.05)])
        y += 1.86
    footer(s, 3)
    return {
        "n": 3,
        "kicker": "Câu hỏi nghiên cứu",
        "title": "Ba câu hỏi, trả lời ngắn trước",
        "seconds": 45,
        "cut": "đủ giờ",
        "goal": "Khóa RQ1–RQ3 và hướng kết luận, để các slide sau là bằng chứng.",
        "say": (
            "Em trả lời ba câu hỏi trong cấu hình không GPU, không Neo4j, và không gọi mô hình sinh câu. "
            "RQ1 hỏi thước tự động khớp người tới đâu. Hướng trả lời: thước có xét số bám người hơn thước đếm từ, "
            "nhưng tập người còn quá nhỏ. "
            "RQ2 hỏi đối chiếu từng khoản, kèm rule số, có đúng hơn gộp cả điều hay không. "
            "Trên 14 mệnh đề đã biết thì có. Trên toàn văn, lớp nói ngược chỉ đúng khoảng 0,15. "
            "RQ3 hỏi vòng sửa có giảm contradicted hay không. Có, khi thay lại đúng khoản. "
            "Vòng này chưa phải mô hình viết lại câu. "
            "Các slide sau là bằng chứng cho ba câu vừa rồi, không phải kết quả bất ngờ."
        ),
        "point": "Dòng 0,15 ở RQ2.",
        "bridge": "Bốn sản phẩm trong mã tương ứng với ba câu hỏi này.",
        "if_asked": "Nếu thầy cô hỏi vì sao kết luận sớm: em muốn buổi nói ngắn, số liệu ở sau là chỗ kiểm lại câu vừa nói.",
    }


def slide_04(prs):
    s = blank(prs)
    header(s, "Đóng góp", "Bốn thứ đã chạy được trong mã")
    items = [
        ("01", TEAL, MINT, "Đồ thị và bộ câu", "Ba thông tư công khai, đồ thị điều–khoản, và ViEduReg trên Thông tư 56/2026."),
        ("02", SKY, SKY_BG, "GGF và ba nhãn", "Supported, contradicted, unsupported, kèm guideline gán nhãn."),
        ("03", GRAPE, GRAPE_BG, "Năm cổng, một chỉ mục", "Cùng BM25, để so cách đi trên đồ thị khi phần tìm từ giống nhau."),
        ("04", CORAL, PEACH, "Chỗ thước chưa đứng", "Điểm gần tuyệt đối trên 14 mệnh đề không giữ được khi dựng mệnh đề từ toàn văn."),
    ]
    positions = [(0.48, 1.22), (6.68, 1.22), (0.48, 4.05), (6.68, 4.05)]
    for (x, y), (num, accent, bg, title, body) in zip(positions, items):
        rrect(s, x, y, 6.15, 2.65, WHITE, 0.1)
        oval(s, x + 0.28, y + 0.32, 0.62, 0.62, accent)
        tb(s, x + 0.28, y + 0.44, 0.62, 0.4, [P(num, 14, WHITE, True, "center")], anchor="middle")
        tb(s, x + 1.1, y + 0.38, 4.7, 0.5, [P(title, 20, INK, True)])
        tb(s, x + 0.28, y + 1.2, 5.6, 1.15, [P(body, 15, MUTED, leading=1.05)])
    footer(s, 4)
    return {
        "n": 4,
        "kicker": "Đóng góp",
        "title": "Bốn thứ đã chạy được trong mã",
        "seconds": 35,
        "cut": "đủ giờ",
        "goal": "Nói rõ bốn mục đã chạy, không phải phần dự kiến.",
        "say": (
            "Bốn thứ này đã chạy trong mã. "
            "Một, đồ thị điều khoản cho ba thông tư, và bộ câu ViEduReg. "
            "Hai, thước GGF cùng ba nhãn và guideline. "
            "Ba, năm cổng truy hồi trên cùng một chỉ mục BM25. "
            "Bốn, và đây là đóng góp em muốn giữ: điểm gần tuyệt đối trên 14 mệnh đề người "
            "không giữ được khi mệnh đề được dựng từ toàn văn. Lớp nói ngược là lớp yếu."
        ),
        "point": "Thẻ thứ tư.",
        "bridge": "Em đi một lượt qua luồng của một câu hỏi.",
    }


def slide_05(prs):
    s = blank(prs)
    header(s, "Phương pháp", "Một câu hỏi đi qua bảy bước")
    steps = [
        ("1", "Hỏi", TEAL),
        ("2", "Truy hồi", TEAL),
        ("3", "Trích khoản", SKY),
        ("4", "Tách mệnh đề", SKY),
        ("5", "Ba nhãn", GRAPE),
        ("6", "Điểm GGF", CORAL),
        ("7", "Sửa ≤ 2 vòng", CORAL),
    ]
    x = 0.42
    for i, (num, label, color) in enumerate(steps):
        rrect(s, x, 1.35, 1.62, 1.7, WHITE, 0.12)
        oval(s, x + 0.52, 1.5, 0.55, 0.55, color)
        tb(s, x + 0.52, 1.6, 0.55, 0.38, [P(num, 14, WHITE, True, "center")], anchor="middle")
        tb(s, x + 0.08, 2.15, 1.46, 0.7, [P(label, 13, INK, True, "center")])
        if i < len(steps) - 1:
            tb(s, x + 1.48, 1.9, 0.28, 0.4, [P("›", 18, TEAL, True, "center")])
        x += 1.84
    notes_cards = [
        (MINT, "Câu trả lời là đoạn trích", "Không gọi mô hình viết lại. Đo việc chọn đúng khoản và bắt số bị đổi."),
        (SKY_BG, "Rule số không bị phủ quyết", "Lệch số hoặc lệch hệ quả thì bộ học được không được gán supported."),
        (PEACH, "Sửa là thay đúng khoản", "Mệnh đề lệch được thay bằng đoạn đang đối chiếu. Tối đa hai vòng."),
    ]
    x = 0.48
    for bg, title, body in notes_cards:
        rrect(s, x, 3.4, 4.0, 2.45, bg, 0.1)
        tb(s, x + 0.22, 3.58, 3.55, 0.7, [P(title, 16, INK, True)])
        tb(s, x + 0.22, 4.35, 3.55, 1.25, [P(body, 14, MUTED, leading=1.05)])
        x += 4.16
    footer(s, 5)
    return {
        "n": 5,
        "kicker": "Phương pháp",
        "title": "Một câu hỏi đi qua bảy bước",
        "seconds": 50,
        "cut": "đủ giờ",
        "goal": "Một lượt luồng. Nhấn câu trả lời là đoạn trích, sửa không gọi LLM.",
        "say": (
            "Một câu hỏi đi từ trái sang phải. "
            "Vào một trong năm cổng truy hồi. Câu trả lời trong bài là đoạn trích khoản, không phải câu mô hình viết. "
            "Em cố ý làm vậy để đo hai việc: có chọn đúng khoản không, và có bắt số bị đổi không. "
            "Đoạn trích được cắt thành mệnh đề, mỗi mệnh đề nhận một trong ba nhãn. "
            "GGF là tỉ lệ mệnh đề được nguồn kéo theo. "
            "Mệnh đề lệch được sửa tối đa hai vòng bằng cách thay lại đúng khoản. "
            "Nếu lệch số hoặc lệch hệ quả, bộ học được không được phủ quyết thành supported. "
            "Máy chạy Python trên CPU. File Cypher để nạp Neo4j đã có, buổi này không cần mở Neo4j."
        ),
        "point": "Ba thẻ dưới: đoạn trích, rule không bị phủ quyết, sửa là thay khoản.",
        "bridge": "Đơn vị của đồ thị và nghĩa của ba nhãn nằm ở slide sau.",
    }


def slide_06(prs):
    s = blank(prs)
    header(s, "Đồ thị và nhãn", "Đơn vị chấm là mệnh đề, đơn vị văn bản là khoản")
    rrect(s, 0.48, 1.2, 6.15, 4.35, WHITE, 0.1)
    tb(s, 0.72, 1.38, 5.7, 0.4, [P("Đồ thị điều–khoản", 18, INK, True)])
    tb(
        s, 0.72, 1.9, 5.7, 1.15,
        [P("Nút: Văn bản → Điều → Khoản → Điểm.\nMỗi điểm giữ đoạn gốc, để từ nút đọc lại đúng câu.", 15, INK, leading=1.05)],
    )
    tb(
        s, 0.72, 3.15, 5.7, 1.15,
        [
            P([R("REFERS_TO", True, TEAL_DEEP, 15), R("  khoản này dẫn khoản kia.", False, INK, 15)], after=3),
            P([R("REPLACES", True, TEAL_DEEP, 15), R("  thông tư này thay thông tư kia.", False, INK, 15)]),
        ],
    )
    tb(
        s, 0.72, 4.4, 5.7, 0.9,
        [P("Toàn văn: 545 đơn vị · 125 dẫn chiếu · 2 thay thế.\n23/2021 hết hiệu lực 30.06.2026. 56/2026 có hiệu lực 07.07.2026.", 13, MUTED, leading=1.05)],
    )
    labels = [
        (TEAL_DEEP, MINT, "Supported", "Khoản kéo theo mệnh đề. Số, ngày và điều kiện khớp nguồn."),
        (CORAL_DEEP, PEACH, "Contradicted", "Nguồn nói ngược. Khác số, khác ngày, hoặc đổi hệ quả. 24 tín chỉ mà nói buộc thôi học vẫn là contradicted."),
        (AMBER_INK, SUN, "Unsupported", "Thêm điều không có trong khoản, và nguồn cũng không nói ngược."),
    ]
    y = 1.2
    for accent, bg, name, body in labels:
        rrect(s, 6.85, y, 6.0, 1.38, bg, 0.1)
        tb(s, 7.08, y + 0.12, 5.55, 0.36, [P(name, 16, accent, True)])
        tb(s, 7.08, y + 0.5, 5.55, 0.75, [P(body, 13, INK, leading=1.02)])
        y += 1.5
    formula = rrect(s, 0.48, 5.7, 12.38, 1.2, TEAL, 0.1)
    fill_tf(
        formula,
        [
            P("GGF  =  số mệnh đề supported  /  tổng số mệnh đề", 20, WHITE, True, "center", after=2),
            P("Tỉ lệ contradicted báo cáo riêng. Không gộp vào một chữ “sai” chung với unsupported.", 13, "E7FFF9", align="center"),
        ],
        anchor="middle", ml=0.2, mr=0.2, mt=0.08, mb=0.06,
    )
    footer(s, 6)
    return {
        "n": 6,
        "kicker": "Đồ thị và nhãn",
        "title": "Đơn vị chấm là mệnh đề, đơn vị văn bản là khoản",
        "seconds": 65,
        "cut": "ngắn",
        "goal": "Khóa schema, hai cạnh, ba nhãn, và công thức GGF. Nhắc hiệu lực hai thông tư.",
        "say": (
            "Đồ thị giữ đúng đơn vị của văn bản quy phạm: văn bản, điều, khoản, điểm. "
            "Hai cạnh rút bằng luật. REFERS_TO khi đoạn dẫn khoản khác. REPLACES khi một thông tư thay thông tư khác. "
            "Toàn văn có 545 khoản hoặc điểm, 125 cạnh dẫn chiếu, hai cạnh thay thế. "
            "Ba thông tư là 08/2021, 23/2021 và 56/2026. "
            "Thông tư 23 hết hiệu lực từ 30 tháng 6 năm 2026. Thông tư 56 có hiệu lực 7 tháng 7 năm 2026, "
            "cho khóa tuyển từ năm 2026. Khóa đã có thông báo tuyển sinh trước ngày đó vẫn theo quy chế cũ. "
            "Ba nhãn dùng chung cho người và cho rule. "
            "Supported là khoản kéo theo mệnh đề. "
            "Contradicted là nguồn nói ngược. Câu vượt quá 24 thì buộc thôi học vẫn là contradicted, "
            "dù số 24 đúng, vì khoản đang đối chiếu chỉ nói cảnh báo. "
            "Unsupported là thêm điều không có trong nguồn. "
            "GGF bằng số mệnh đề supported chia cho tổng số mệnh đề. Contradicted em báo cáo riêng."
        ),
        "point": "Công thức ở dải xanh, rồi ví dụ 24 tín chỉ trên thẻ contradicted.",
        "bridge": "Năm cổng truy hồi cùng đọc đồ thị này, nhưng không phải mã gốc của các bài em dẫn.",
    }


def slide_07(prs):
    s = blank(prs)
    header(s, "Truy hồi", "Năm cổng, một chỉ mục BM25")
    rows = [
        ("Naive", "Lấy đoạn BM25 đứng đầu."),
        ("Graph", "Lấy hạt giống đó, rồi đi cạnh dẫn chiếu."),
        ("LightRAG-style", "Trộn đoạn khớp nội dung với đoạn có tiêu đề điều khớp câu hỏi."),
        ("HippoRAG-style", "Ba hạt giống, cộng điểm cho hàng xóm và đoạn cùng điều. Lan một lớp cạnh."),
        ("KAG-style", "Ưu tiên đoạn nói hiệu lực hoặc điều cấm khi câu hỏi hỏi việc đó, rồi mới BM25."),
    ]
    y = 1.2
    for i, (name, desc) in enumerate(rows):
        bg = MINT if i == 3 else WHITE
        rrect(s, 0.48, y, 12.38, 0.92, bg, 0.12)
        tb(s, 0.72, y + 0.22, 2.7, 0.5, [P(name, 16, TEAL_DEEP if i == 3 else INK, True)], anchor="middle")
        tb(s, 3.5, y + 0.16, 9.0, 0.62, [P(desc, 15, INK)], anchor="middle", ml=0.08)
        y += 1.02
    footer(s, 7)
    return {
        "n": 7,
        "kicker": "Truy hồi",
        "title": "Năm cổng, một chỉ mục BM25",
        "seconds": 40,
        "cut": "đủ giờ",
        "goal": "Nói khác nhau ở cách dùng đồ thị. Chặn hiểu nhầm là đã chạy mã gốc.",
        "say": (
            "Năm cổng đọc cùng chỉ mục BM25, đã chuẩn hóa dấu và viết thường. "
            "Lần này không dùng embedding dày, vì em muốn so cách đi trên đồ thị khi phần tìm từ giống nhau. "
            "Naive lấy đoạn đứng đầu. Graph đi thêm cạnh dẫn chiếu từ một hạt giống. "
            "LightRAG-style trộn nội dung với tiêu đề điều. "
            "HippoRAG-style lấy ba hạt giống rồi lan một lớp cạnh. Đây không phải OpenIE bằng mô hình lớn. "
            "KAG-style ưu tiên đoạn hiệu lực hoặc điều cấm khi câu hỏi hỏi những việc đó. "
            "Em không chạy Docker gốc của LightRAG, HippoRAG hay KAG. Tên trên slide là loại thuật toán đã cổng hóa."
        ),
        "point": "Dòng HippoRAG-style, vì slide kết quả sẽ nhắc cổng này.",
        "bridge": "Trước khi đọc Recall, em tách ba tầng dữ liệu để các số không bị trộn.",
        "if_asked": "Nếu bị hỏi có phải HippoRAG 2 không: không. Cùng ý tưởng lan từ hạt giống, trên BM25, một lớp cạnh.",
    }


def slide_08(prs):
    s = blank(prs)
    header(s, "Dữ liệu", "Ba tầng dữ liệu, không trộn khi đọc số")
    stats = [
        (TEAL, "545", "khoản hoặc điểm", "Đồ thị toàn văn ba thông tư.\n125 cạnh dẫn chiếu, 2 cạnh thay thế."),
        (SKY, "150", "câu đã đối chiếu", "ViEduReg trên 56/2026.\n90 câu một khoản, 60 câu nhiều khoản."),
        (CORAL, "14", "mệnh đề người", "Test trên trích Thông tư 08.\nĐủ để kiểm lỗi guideline, chưa đủ để suy tương quan."),
    ]
    x = 0.48
    for accent, num, label, body in stats:
        rrect(s, x, 1.22, 4.0, 3.15, WHITE, 0.1)
        tb(s, x + 0.22, 1.4, 3.55, 0.85, [P(num, 40, accent, True)])
        tb(s, x + 0.22, 2.25, 3.55, 0.4, [P(label, 16, INK, True)])
        tb(s, x + 0.22, 2.8, 3.55, 1.3, [P(body, 14, MUTED, leading=1.05)])
        x += 4.16
    rrect(s, 0.48, 4.55, 12.38, 2.3, WHITE, 0.1)
    tb(s, 0.74, 4.72, 8, 0.4, [P("127 mệnh đề không trùng, một lượt đối chiếu", 16, INK, True)])
    chips = [("72", "supported", TEAL_DEEP, MINT), ("26", "contradicted", CORAL_DEEP, PEACH), ("29", "unsupported", AMBER_INK, SUN)]
    x = 0.74
    for num, name, accent, bg in chips:
        rrect(s, x, 5.25, 2.5, 1.25, bg, 0.12)
        tb(s, x + 0.15, 5.35, 2.2, 0.55, [P(num, 26, accent, True, "center")])
        tb(s, x + 0.15, 5.9, 2.2, 0.4, [P(name, 13, INK, True, "center")])
        x += 2.7
    tb(
        s, 8.9, 5.2, 3.7, 1.35,
        [P("Chưa có người thứ hai,\nnên chưa có kappa.\nNhãn này không phải macro-F1 ở các slide sau.", 14, MUTED, leading=1.05)],
    )
    footer(s, 8)
    return {
        "n": 8,
        "kicker": "Dữ liệu",
        "title": "Ba tầng dữ liệu, không trộn khi đọc số",
        "seconds": 45,
        "cut": "đủ giờ",
        "goal": "Tách 545 / 150 / 14. Nhấn 14 mới là test người, 127 chưa phải macro-F1.",
        "say": (
            "Em tách ba tầng, và không trộn khi đọc số. "
            "Tầng rộng là 545 đơn vị trên ba thông tư, thu ngày 30 tháng 10. "
            "Thông tư 08 lấy từ PDF Công báo. Thông tư 56 là toàn văn công khai, không phải file Công báo gốc. "
            "Tầng câu hỏi là 150 câu ViEduReg đã đối chiếu nguồn. "
            "Từ đầu ra các hệ, em gom 127 mệnh đề không trùng. Một lượt đối chiếu cho 72 supported, 26 contradicted, 29 unsupported. "
            "Chưa có người thứ hai, nên chưa có kappa. Các số macro-F1 ở sau không tính trên 127 mệnh đề này. "
            "Tầng test người là 14 mệnh đề trên đồ thị trích của Thông tư 08. "
            "Tập này đủ để kiểm lỗi đã viết trong guideline, và quá nhỏ để ước lượng tương quan. "
            "Huấn luyện verifier chủ yếu là mệnh đề tổng hợp, 2.688 train và 298 dev. Tập người không trộn vào huấn luyện."
        ),
        "point": "Ba số lớn 545, 150, 14, rồi dòng chưa có kappa.",
        "bridge": "Slide truy hồi dùng 150 câu. Slide verifier dùng 14 mệnh đề và mệnh đề dựng từ toàn văn.",
    }


def slide_09(prs):
    s = blank(prs)
    header(s, "Kết quả truy hồi", "Tìm được khoản chưa có nghĩa đã đưa vào câu")
    tb(s, 0.5, 1.08, 8.2, 0.28, [P("150 câu ViEduReg. Thanh đậm: Recall pool. Thanh nhạt: Recall câu trả lời.", 13, MUTED)])
    systems = [
        ("Naive", 0.440, 0.200, False),
        ("Graph", 0.400, 0.257, False),
        ("LightRAG-style", 0.413, 0.260, False),
        ("HippoRAG-style", 0.633, 0.270, True),
        ("KAG-style", 0.460, 0.213, False),
    ]
    y = 1.42
    for name, pool, ans, hi in systems:
        if hi:
            rrect(s, 0.42, y - 0.04, 8.35, 0.84, MINT, 0.1)
        tb(s, 0.5, y, 2.15, 0.74, [P(name, 13, TEAL_DEEP if hi else INK, True)], anchor="middle")
        bar(s, 2.7, y + 0.06, 4.55, 0.26, pool, TEAL)
        bar(s, 2.7, y + 0.4, 4.55, 0.26, ans, "FFB0A6")
        tb(s, 7.35, y, 1.3, 0.34, [P(f"{pool:.3f}".replace(".", ","), 12, TEAL_DEEP, True)], anchor="middle")
        tb(s, 7.35, y + 0.34, 1.3, 0.34, [P(f"{ans:.3f}".replace(".", ","), 12, CORAL_DEEP, True)], anchor="middle")
        y += 0.88
    rrect(s, 8.95, 1.42, 3.9, 4.28, WHITE, 0.1)
    tb(s, 9.15, 1.56, 3.5, 0.4, [P("Đọc hai thanh", 16, INK, True)])
    tb(
        s, 9.15, 2.08, 3.5, 3.4,
        [
            P("Pool HippoRAG-style đạt 0,633.", 14, INK, after=6, leading=1.02),
            P("Câu trả lời của mọi hệ chỉ khoảng 0,20 đến 0,27.", 14, INK, after=6, leading=1.02),
            P("Câu hai bước: pool gần đủ, 0,78–0,98. Câu trả lời chỉ khoảng 0,45–0,60.", 14, INK, after=6, leading=1.02),
            P("Verifier phải nhìn pool, không chỉ nhìn câu đã ghép.", 14, TEAL_DEEP, True, leading=1.02),
        ],
    )
    chips = [
        ("0,563", "HippoRAG-style, Recall@5 trên 63 câu toàn văn"),
        ("0,421", "Naive, Recall@5 trên cùng 63 câu"),
        ("0,625", "HippoRAG-style trên câu viện dẫn, naive là 0,458"),
    ]
    x = 0.48
    for num, label in chips:
        rrect(s, x, 5.88, 4.02, 1.08, WHITE, 0.12)
        tb(s, x + 0.14, 5.98, 1.15, 0.86, [P(num, 18, TEAL_DEEP, True, "center")], anchor="middle")
        tb(s, x + 1.3, 6.02, 2.55, 0.8, [P(label, 12, MUTED, leading=1.0)], anchor="middle")
        x += 4.16
    footer(s, 9)
    return {
        "n": 9,
        "kicker": "Kết quả truy hồi",
        "title": "Tìm được khoản chưa có nghĩa đã đưa vào câu",
        "seconds": 70,
        "cut": "ngắn",
        "goal": "Tách Recall pool và Recall câu trả lời. Nhắc lợi thế viện dẫn và rủi ro một hạt giống sai.",
        "say": (
            "Slide này tách hai việc hay bị gộp: khoản vàng có nằm trong pool hay không, "
            "và khoản vàng có được viết vào câu trả lời hay không. "
            "Thanh đậm là pool. Thanh nhạt là câu trả lời. "
            "HippoRAG-style đứng đầu pool, 0,633. Naive pool là 0,440. "
            "Khi chỉ tính khoản ghép vào câu, mọi hệ nằm khoảng 0,20 đến 0,27. "
            "Tìm được khoản chưa có nghĩa đã đưa khoản vào câu. "
            "Với câu hai bước, pool gần như đủ, từ 0,78 đến 0,98, nhưng câu trả lời chỉ phủ khoảng một nửa đến 0,60. "
            "Verifier vì thế phải nhìn pool. "
            "Trên đồ thị toàn văn 63 câu, Recall@5 của HippoRAG-style là 0,563, naive là 0,421. "
            "Lợi thế rõ ở câu viện dẫn: 0,625 so với 0,458. "
            "Graph đi từ một hạt giống chỉ đạt 0,333 chung, vì hạt giống sai thì cạnh đi ra cũng sai."
        ),
        "point": "Hai thanh của dòng HippoRAG-style, rồi ba số nhỏ ở chân slide.",
        "bridge": "Truy hồi là chuyện chọn đoạn. Slide sau là chuyện nhận ra đoạn đó có kéo theo mệnh đề hay không.",
    }


def slide_10(prs):
    s = blank(prs)
    header(s, "Verifier  ·  RQ2", "Điểm 1,000 không sống được trên toàn văn")
    tb(s, 0.5, 1.12, 8, 0.3, [P("Macro-F1 trên 14 mệnh đề người của Thông tư 08/2021.", 13, MUTED)])
    rows = [
        ("Từng khoản + rule số", 1.000, TEAL, TEAL_DEEP),
        ("Gộp cả điều", 0.941, SKY, "1D4ED8"),
        ("Bỏ rule số", 0.436, CORAL, CORAL_DEEP),
        ("TF-IDF + logistic", 0.578, AMBER, AMBER_INK),
        ("PhoBERT-base, 1 epoch", 0.419, GRAPE, "4338CA"),
    ]
    y = 1.5
    for name, val, bar_color, text_color in rows:
        tb(s, 0.48, y, 3.15, 0.7, [P(name, 13, INK, True)], anchor="middle")
        bar(s, 3.7, y + 0.2, 4.3, 0.32, val, bar_color)
        tb(s, 8.1, y, 0.95, 0.7, [P(f"{val:.3f}".replace(".", ","), 14, text_color, True)], anchor="middle")
        y += 0.78
    card = rrect(s, 9.15, 1.48, 3.7, 2.55, CORAL_DEEP, 0.1)
    fill_tf(
        card,
        [
            P("0,15", 48, WHITE, True, "center", after=0),
            P("độ đúng lớp contradicted\ntrên mệnh đề dựng từ toàn văn", 13, "FFF6F4", align="center", leading=1.02),
        ],
        anchor="middle", ml=0.16, mr=0.14, mt=0.08, mb=0.08,
    )
    rrect(s, 9.15, 4.18, 3.7, 1.85, WHITE, 0.1)
    tb(s, 9.35, 4.32, 3.35, 1.55, [
        P("PhoBERT trên dev tổng hợp", 13, MUTED, after=2),
        P("0,925  →  0,419", 22, GRAPE, True, after=2),
        P("Dev cao, 14 mệnh đề người thì thấp hơn TF-IDF.", 12, MUTED, leading=1.02),
    ])
    note = rrect(s, 0.48, 5.55, 8.5, 1.35, SUN, 0.1)
    fill_tf(
        note,
        [P(
            "Điểm 1,000 là kiểm tra trên lỗi guideline đã viết rule để bắt. "
            "Một mệnh đề chỉ đúng khi nhìn từng khoản: nợ 24 tín chỉ thì buộc thôi học. "
            "Gộp cả điều thì số và hệ quả cùng nằm trong túi từ, nên bị gán supported.",
            13, INK, leading=1.05,
        )],
        ml=0.18, mr=0.16, mt=0.12, mb=0.08,
    )
    footer(s, 10)
    return {
        "n": 10,
        "kicker": "Verifier · RQ2",
        "title": "Điểm 1,000 không sống được trên toàn văn",
        "seconds": 80,
        "cut": "ngắn",
        "goal": "Macro-F1 1,000 là cận trên. Contradicted toàn văn 0,15 là kết quả chính. PhoBERT không vượt TF-IDF.",
        "say": (
            "Đây là slide em muốn thầy cô nhìn kỹ. "
            "Trên 14 mệnh đề người, rule từng khoản đạt macro-F1 bằng 1. Gộp cả điều còn 0,941. Bỏ rule số thì còn 0,436. "
            "Có một mệnh đề chỉ đúng khi nhìn từng khoản: nợ vượt quá 24 thì bị buộc thôi học. "
            "Số 24 nằm ở điểm cảnh báo. Chữ buộc thôi học nằm ở khoản khác cùng điều. "
            "Gộp cả điều thì cả số lẫn hệ quả đều có trong túi từ, nên bị gán supported. "
            "Điểm 1,000 là cận trên trên các lỗi đã viết trong guideline. "
            "Khi dựng mệnh đề từ khoảng 80 khoản của toàn văn, supported và unsupported vẫn gần đúng, "
            "nhưng contradicted chỉ đúng 0,15. Nhiều số bị đổi không nằm trong đơn vị rule biết so. "
            "Độ phủ từ còn cao nên rule gán supported. "
            "TF-IDF đạt 0,578 trên 14 mệnh đề người. PhoBERT-base, một epoch, đạt 0,925 trên dev tổng hợp "
            "nhưng chỉ 0,419 trên người, thấp hơn TF-IDF. F1 contradicted của PhoBERT là 0,222."
        ),
        "point": "Số 0,15 trên thẻ đỏ, rồi mũi tên 0,925 sang 0,419.",
        "bridge": "RQ2 là vậy. Em gộp RQ1 và vòng sửa ở slide kế, vì cả hai cần đọc cùng hạn chế của tập nhỏ.",
    }


def slide_11(prs):
    s = blank(prs)
    header(s, "RQ1 và RQ3", "Khớp người còn mỏng. Sửa xóa được lỗi đã chỉ đúng")
    rrect(s, 0.48, 1.2, 6.15, 4.55, WHITE, 0.1)
    tb(s, 0.72, 1.38, 5.7, 0.4, [P("RQ1  ·  khớp với người", 16, TEAL_DEEP, True)])
    pairs = [
        ("GGF rule, từng khoản", "1,000", "Rất ít nhóm câu"),
        ("Thước đếm từ", "−0,052", "Câu sai số vẫn gần nguyên văn"),
        ("RAGAS-style, gộp điều", "0,750", "Không phải RAGAS gốc"),
        ("Giám khảo = GGF rule", "0,763", "Không phải mô hình lớn"),
    ]
    y = 1.95
    for name, val, note in pairs:
        tb(s, 0.75, y, 3.35, 0.32, [P(name, 14, INK, True)])
        tb(s, 4.2, y, 2.1, 0.34, [P(val, 20, TEAL_DEEP, True, "right")])
        tb(s, 0.75, y + 0.36, 5.45, 0.3, [P(note, 12, MUTED)])
        y += 0.86
    rrect(s, 6.85, 1.2, 6.0, 4.55, WHITE, 0.1)
    tb(s, 7.08, 1.38, 5.55, 0.4, [P("RQ3  ·  vòng sửa tất định", 16, CORAL_DEEP, True)])
    blocks = [
        (PEACH, "13 câu cố ý đổi mốc", "GGF  0,231  →  1", "Contradicted  0,744  →  0"),
        (SUN, "20 câu đổi số, toàn văn", "GGF  0,45  →  1", "Contradicted  0,36  →  0"),
    ]
    y = 1.95
    for bg, title, a, b in blocks:
        rrect(s, 7.08, y, 5.55, 1.45, bg, 0.1)
        tb(s, 7.25, y + 0.1, 5.2, 0.32, [P(title, 13, MUTED)])
        tb(s, 7.25, y + 0.42, 5.2, 0.4, [P(a, 18, INK, True)])
        tb(s, 7.25, y + 0.9, 5.2, 0.38, [P(b, 14, CORAL_DEEP, True)])
        y += 1.58
    tb(s, 7.08, 5.15, 5.55, 0.45, [P("Cơ chế: thay mảnh lệch bằng đúng khoản. Không gọi mô hình viết lại.", 13, MUTED)])
    foot = rrect(s, 0.48, 5.9, 12.38, 1.0, MINT, 0.1)
    fill_tf(
        foot,
        [P("Chưa kết luận tương quan ổn định. Vòng sửa chứng minh: khi verifier chỉ đúng khoản, thay lại đoạn thì contradicted hết. Chưa đo F1 của câu do mô hình viết.", 15, TEAL_DEEP, leading=1.05)],
        anchor="middle", ml=0.24, mr=0.2, mt=0.08, mb=0.06,
    )
    footer(s, 11)
    return {
        "n": 11,
        "kicker": "RQ1 và RQ3",
        "title": "Khớp người còn mỏng. Sửa xóa được lỗi đã chỉ đúng",
        "seconds": 65,
        "cut": "ngắn",
        "goal": "RQ1: đừng gọi là LLM-judge. RQ3: GGF tăng vì thay đúng khoản.",
        "say": (
            "Bên trái là RQ1. Spearman giữa GGF của rule từng khoản và nhãn người bằng 1, trên rất ít nhóm. "
            "Thước đếm từ gần như không tương quan, khoảng âm 0,05, vì câu sai số vẫn gần nguyên văn khoản. "
            "RAGAS-style trong bài là 0,75. Giám khảo trong bài là 0,76, và giám khảo đó chính là GGF của rule, không phải mô hình lớn. "
            "Em không gọi RAGAS gốc. Với vài nhóm, em chưa kết luận tương quan ổn định. "
            "Điều hữu ích là thước có xét số bám người hơn thước đếm từ. "
            "Bên phải là RQ3. Trên 13 câu cố ý đổi mốc, GGF từ 0,231 lên 1, contradicted từ 0,744 về 0. "
            "Trên 20 câu đổi số của toàn văn, GGF từ 0,45 lên 1. "
            "Cơ chế là thay mảnh lệch bằng đoạn khoản đang đối chiếu. "
            "Vòng này chứng minh khi verifier chỉ đúng khoản thì sửa tất định xóa được contradicted. "
            "Chưa chứng minh một mô hình sinh viết lại câu cho đúng."
        ),
        "point": "Cột phải, hai mũi tên về 1. Rồi dải xanh dưới: chưa kết luận tương quan.",
        "bridge": "Em mở demo ba nhịp cho đúng ba ý vừa nói: câu đúng, câu nói ngược, và vòng sửa.",
    }


def slide_12(prs):
    s = blank(prs)
    header(s, "Demo", "Ba nhịp trên trang local")
    cards = [
        (TEAL, MINT, "1", "Câu đúng", "Hỏi nợ bao nhiêu tín chỉ thì bị cảnh báo.",
         "Điều 11 khoản 1 điểm a. Có cụm vượt quá 24. Nhãn supported. GGF bằng 1."),
        (CORAL, PEACH, "2", "Nói ngược", "Đưa câu: vượt quá 24 thì bị buộc thôi học.",
         "Nhãn contradicted. Gộp cả điều có thể thành supported, vì hai ý đứng cạnh nhau."),
        (AMBER, SUN, "3", "Sửa số", "Đổi thời gian tối đa thành 03 lần.",
         "Khoản ghi 02 lần. Sau một vòng, câu về 02 và GGF về 1."),
    ]
    x = 0.48
    for accent, bg, num, title, a, b in cards:
        rrect(s, x, 1.22, 4.0, 4.35, bg, 0.1)
        oval(s, x + 0.24, 1.44, 0.55, 0.55, accent)
        tb(s, x + 0.24, 1.54, 0.55, 0.38, [P(num, 16, WHITE, True, "center")], anchor="middle")
        tb(s, x + 0.95, 1.5, 2.8, 0.48, [P(title, 20, INK, True)])
        tb(s, x + 0.24, 2.25, 3.52, 1.15, [P(a, 15, INK, leading=1.05)])
        tb(s, x + 0.24, 3.5, 3.52, 1.75, [P(b, 14, MUTED, leading=1.05)])
        x += 4.16
    bar_ = rrect(s, 0.48, 5.75, 12.38, 1.15, INK, 0.1)
    fill_tf(
        bar_,
        [
            P("python run.py web", 18, WHITE, True, "center", after=1),
            P("http://127.0.0.1:8766     ·     nếu không mở được trình duyệt: python run.py demo", 13, "F6E7DF", align="center"),
        ],
        anchor="middle", ml=0.2, mr=0.2, mt=0.08, mb=0.06,
    )
    footer(s, 12)
    return {
        "n": 12,
        "kicker": "Demo",
        "title": "Ba nhịp trên trang local",
        "seconds": 30,
        "cut": "demo",
        "goal": "Không đọc slide. Làm trên web. Slide này là phao nếu quên câu.",
        "say": (
            "Em chuyển sang trang local, cổng 8766. Slide này chỉ là phao. "
            "Nhịp một, hỏi nợ bao nhiêu tín chỉ thì bị cảnh báo. Kỳ vọng trích Điều 11, có vượt quá 24, nhãn supported, GGF bằng 1. "
            "Nhịp hai, đưa câu vượt quá 24 thì bị buộc thôi học. Kỳ vọng contradicted. "
            "Nếu gộp cả điều, nhãn có thể thành supported. Đó là lý do phải chấm từng khoản. "
            "Nhịp ba, câu thời gian tối đa bị đổi thành 03 lần, trong khi khoản ghi 02. Sau một vòng sửa, GGF về 1. "
            "Nếu trình duyệt không mở, em chạy python run.py demo, cùng ba ý này."
        ),
        "point": "Làm trên trình duyệt, liếc slide khi cần câu mẫu.",
        "bridge": "Sau demo, em nói chỗ chưa được kết luận, rồi mới kết.",
    }


def slide_13(prs):
    s = blank(prs)
    header(s, "Hạn chế", "Những chỗ chưa được kết luận")
    items = [
        ("Chưa có kappa", "Chưa có người chấm thứ hai cho câu hỏi và cho mệnh đề."),
        ("14 mệnh đề là quá nhỏ", "Đủ để kiểm lỗi đã biết. Chưa đủ để suy tương quan với người."),
        ("ViHERMES chưa vào bảng", "Có trong demo. Bộ trên Hugging Face cần tài khoản."),
        ("Không phải mã gốc", "Năm cổng là thuật toán trên cùng BM25, không phải Docker của bài được dẫn."),
        ("56/2026 không phải Công báo", "Toàn văn công khai. Thuộc tính cổng Chính phủ xác nhận số hiệu và ngày."),
        ("Sửa chưa có mô hình sinh", "Chưa đo F1 của câu do mô hình viết sau khi sửa."),
    ]
    positions = []
    for r in range(2):
        for c in range(3):
            positions.append((0.48 + c * 4.16, 1.22 + r * 2.55))
    for (x, y), (title, body) in zip(positions, items):
        rrect(s, x, y, 4.0, 2.4, WHITE, 0.1)
        rect(s, x, y, 0.1, 2.4, CORAL if "kappa" in title or "nhỏ" in title or "sinh" in title else TEAL)
        tb(s, x + 0.28, y + 0.25, 3.5, 0.7, [P(title, 16, INK, True)])
        tb(s, x + 0.28, y + 1.05, 3.5, 1.1, [P(body, 14, MUTED, leading=1.05)])
    footer(s, 13)
    return {
        "n": 13,
        "kicker": "Hạn chế",
        "title": "Những chỗ chưa được kết luận",
        "seconds": 40,
        "cut": "ngắn",
        "goal": "Nói trước khi kết luận, để các bảng không bị đọc quá mức.",
        "say": (
            "Em nêu hạn chế trước khi kết, vì các bảng dễ bị đọc quá mức. "
            "Chưa có người thứ hai, nên chưa có kappa. "
            "Mười bốn mệnh đề quá nhỏ để suy tương quan với người nói chung. "
            "ViHERMES có trong demo nhưng chưa vào bảng, vì bộ cần tài khoản. "
            "Năm hệ là cổng trên cùng BM25, không phải mã gốc. "
            "Thông tư 56 không phải file Công báo. "
            "Câu hỏi toàn văn phần lớn sinh từ chính đoạn vàng, nên Recall không đo câu người dùng hỏi tự do. "
            "Vòng sửa chưa có mô hình sinh."
        ),
        "point": "Thẻ chưa có kappa và thẻ 14 mệnh đề.",
        "bridge": "Trong giới hạn vừa nói, em giữ lại ba điều.",
    }


def slide_14(prs):
    s = blank(prs)
    header(s, "Kết luận", "Ba điều em giữ lại")
    cards = [
        (TEAL, "0,27", "Đưa vào câu", "Pool cao nhất 0,633.\nCâu trả lời chỉ khoảng\n0,20 đến 0,27."),
        (CORAL, "0,15", "Nói ngược, toàn văn", "Rule đạt 1,000 trên 14 mệnh đề.\nBỏ rule số còn 0,436.\nToàn văn chỉ đúng khoảng 0,15."),
        (SKY, "→ 1", "Sửa khi chỉ đúng khoản", "Contradicted về 0\nkhi thay lại đúng đoạn.\nChưa phải câu mô hình viết."),
    ]
    x = 0.48
    for accent, big, title, body in cards:
        rrect(s, x, 1.22, 4.0, 4.15, WHITE, 0.1)
        rect(s, x, 1.22, 4.0, 0.1, accent)
        tb(s, x + 0.2, 1.5, 3.6, 0.9, [P(big, 40, accent, True)])
        tb(s, x + 0.2, 2.5, 3.6, 0.45, [P(title, 16, INK, True)])
        tb(s, x + 0.2, 3.15, 3.6, 1.9, [P(body, 15, MUTED, leading=1.08)])
        x += 4.16
    foot = rrect(s, 0.48, 5.55, 12.38, 1.35, TEAL, 0.1)
    fill_tf(
        foot,
        [
            P("Việc còn lại, trước khi kết luận độ khớp với người", 14, "D7FFF6", align="center", after=2),
            P("Người chấm thứ hai  ·  miền y tế trong bảng  ·  vòng sửa có mô hình sinh", 16, WHITE, True, "center"),
        ],
        anchor="middle", ml=0.2, mr=0.2, mt=0.1, mb=0.08,
    )
    footer(s, 14)
    return {
        "n": 14,
        "kicker": "Kết luận",
        "title": "Ba điều em giữ lại",
        "seconds": 45,
        "cut": "ngắn",
        "goal": "Nhắc lại ba số. Không thêm kết luận mới.",
        "say": (
            "Em giữ lại ba điều. "
            "Một, tìm được khoản trong pool chưa có nghĩa khoản đã vào câu. Pool cao nhất 0,633, câu trả lời chỉ khoảng 0,20 đến 0,27. "
            "Hai, rule số và hệ quả rất mạnh trên lỗi đã khai báo, macro-F1 bằng 1, và còn 0,436 nếu bỏ rule số. "
            "Trên mệnh đề dựng từ toàn văn, cùng rule chỉ đúng khoảng 0,15 mệnh đề contradicted. PhoBERT tinh chỉnh ngắn không vượt TF-IDF. "
            "Ba, vòng sửa tất định đưa contradicted về 0 khi đã chỉ đúng khoản. Đây là thay đoạn, chưa phải mô hình viết lại. "
            "Trước khi kết luận độ khớp với người, cần người chấm thứ hai, đưa miền y tế vào bảng, và một vòng sửa có mô hình sinh."
        ),
        "point": "Ba số 0,27 · 0,15 · mũi tên về 1, rồi dải xanh việc còn lại.",
        "bridge": "Em dừng ở đây và mời câu hỏi. Demo vẫn mở nếu thầy cô muốn xem một câu.",
    }


def slide_15(prs):
    s = blank(prs)
    rect(s, 0, 0, W, H, BG)
    rrect(s, 0.45, 0.4, 12.45, 6.7, TEAL, 0.06)
    oval(s, 10.6, 0.7, 1.8, 1.8, "13B39A")
    oval(s, 0.75, 5.15, 1.3, 1.3, "0C7C70")
    tb(s, 1.1, 1.15, 11, 0.4, [P("HỎI ĐÁP", 14, "D7FFF6", True, "center")])
    tb(s, 1.1, 1.7, 11, 1.1, [P("Cảm ơn thầy cô", 48, WHITE, True, "center")])
    tb(
        s, 1.6, 3.05, 10.1, 0.8,
        [P("Em sẵn sàng mở lại demo, hoặc đi vào một số trên slide.", 18, "F3FFFC", align="center")],
    )
    chips = [
        ("0,633 và 0,27", "pool và câu trả lời"),
        ("1,000 → 0,436", "bỏ rule số"),
        ("0,15", "contradicted toàn văn"),
    ]
    x = 1.15
    for big, small in chips:
        rrect(s, x, 4.15, 3.5, 1.35, "0C7C70", 0.12)
        tb(s, x + 0.1, 4.28, 3.3, 0.55, [P(big, 18, WHITE, True, "center")])
        tb(s, x + 0.1, 4.85, 3.3, 0.4, [P(small, 13, "D7FFF6", align="center")])
        x += 3.7
    tb(s, 1.1, 5.75, 11, 0.45, [P("ViKAG-Faith   ·   30.10.2026", 14, "E7FFF9", align="center")])
    return {
        "n": 15,
        "kicker": "Kết",
        "title": "Cảm ơn thầy cô",
        "seconds": 20,
        "cut": "ngắn",
        "goal": "Dừng. Mời hỏi. Nhắc ba chỗ nếu thầy cô muốn đi sâu.",
        "say": (
            "Em cảm ơn thầy cô đã nghe. Em sẵn sàng mở demo hoặc trả lời câu hỏi. "
            "Nếu thầy cô muốn xem một số cụ thể, em ưu tiên ba chỗ: "
            "Recall pool 0,633 so với câu trả lời khoảng 0,27, "
            "macro-F1 từ 1,000 xuống 0,436 khi bỏ rule số, "
            "và độ đúng 0,15 của lớp contradicted trên toàn văn."
        ),
        "point": "Ba thẻ số ở dưới.",
        "bridge": "",
    }


BUILDERS = [
    slide_01, slide_02, slide_03, slide_04, slide_05,
    slide_06, slide_07, slide_08, slide_09, slide_10,
    slide_11, slide_12, slide_13, slide_14, slide_15,
]


def build_pptx():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    prs.core_properties.title = "ViKAG-Faith — Báo cáo cuối kỳ"
    prs.core_properties.subject = "Xử lý ngôn ngữ tự nhiên"
    prs.core_properties.category = "Thuyết trình"
    prs.core_properties.author = "ViKAG-Faith"
    sections = []
    for fn in BUILDERS:
        sections.append(fn(prs))
    for slide, sec in zip(prs.slides, sections):
        notes(slide, sec)
    prs.save(PPTX_PATH)
    return sections


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

def shade_cell(cell, fill):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag == dqn("w:shd"):
            tcPr.remove(child)
    shd = OxmlElement("w:shd")
    shd.set(dqn("w:val"), "clear")
    shd.set(dqn("w:color"), "auto")
    shd.set(dqn("w:fill"), fill)
    tcPr.append(shd)


def set_cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for m, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{m}")
        node.set(dqn("w:w"), str(val))
        node.set(dqn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def no_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(dqn("w:val"), "nil")
        el.set(dqn("w:sz"), "0")
        el.set(dqn("w:space"), "0")
        el.set(dqn("w:color"), "auto")
        borders.append(el)
    tblPr.append(borders)


def drun(p, text, size=12, bold=False, color="1C2430", italic=False):
    run = p.add_run(text)
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(dqn("w:eastAsia"), "Arial")
    run.font.size = DPt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = DocRGB(int(color[0:2], 16), int(color[2:4], 16), int(color[4:6], 16))
    return run


def dpara(doc, text, size=12, bold=False, color="1C2430", space_after=8, space_before=0, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = DPt(space_after)
    p.paragraph_format.space_before = DPt(space_before)
    p.paragraph_format.line_spacing = 1.12
    drun(p, text, size, bold, color, italic)
    return p


def callout(doc, label, text, fill, label_color="0B6E62"):
    table = doc.add_table(rows=1, cols=1)
    no_table_borders(table)
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    set_cell_margins(cell, 80, 80, 140, 140)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = DPt(0)
    p.paragraph_format.line_spacing = 1.12
    drun(p, label + "  ", 12, True, label_color)
    drun(p, text, 12, False, "1C2430")
    doc.add_paragraph().paragraph_format.space_after = DPt(4)


def set_doc_defaults(doc):
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = DPt(12)
    style.element.rPr.rFonts.set(dqn("w:eastAsia"), "Arial")
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.paragraph_format.space_after = DPt(2)
    drun(hp, "Kịch bản thuyết trình  ·  ViKAG-Faith  ·  Xử lý ngôn ngữ tự nhiên", 10, False, "5C6B7A")
    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    drun(fp, "ViKAG-Faith   ·   trang ", 10, False, "5C6B7A")
    run = fp.add_run()
    run.font.name = "Arial"
    run.font.size = DPt(10)
    run.font.color.rgb = DocRGB(0x5C, 0x6B, 0x7A)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(dqn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(dqn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(dqn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def add_timing_table(doc, sections):
    table = doc.add_table(rows=1, cols=4)
    hdr = ["Slide", "Tiêu đề", "Giây", "Nhịp 12 phút"]
    for i, text in enumerate(hdr):
        cell = table.rows[0].cells[i]
        shade_cell(cell, "0E9A86")
        set_cell_margins(cell, 50, 50, 70, 70)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = DPt(0)
        drun(p, text, 11, True, "FFFFFF")
    for sec in sections:
        row = table.add_row().cells
        vals = [
            f"{sec['n']:02d}",
            sec["title"],
            str(sec["seconds"]),
            "Có" if sec["cut"] == "ngắn" else ("Demo" if sec["cut"] == "demo" else "Nếu còn giờ"),
        ]
        fill = "E7F7F3" if sec["cut"] == "ngắn" else "FFFFFF"
        for i, val in enumerate(vals):
            shade_cell(row[i], fill)
            set_cell_margins(row[i], 40, 40, 70, 70)
            p = row[i].paragraphs[0]
            p.paragraph_format.space_after = DPt(0)
            drun(p, val, 10, i == 0, "1C2430")
    doc.add_paragraph().paragraph_format.space_after = DPt(6)


def build_docx(sections):
    doc = Document()
    set_doc_defaults(doc)
    dpara(doc, "KỊCH BẢN THUYẾT TRÌNH", 12, True, "0E9A86", space_after=2)
    dpara(doc, "ViKAG-Faith", 26, True, "1C2430", space_after=4)
    dpara(
        doc,
        "Đánh giá tính trung thực mức mệnh đề cho hỏi đáp quy chế đào tạo tiếng Việt.",
        13, False, "5C6B7A", space_after=8,
    )
    dpara(
        doc,
        "Môn Xử lý ngôn ngữ tự nhiên  ·  30 tháng 10 năm 2026  ·  đi kèm file ViKAG-Faith_bao_cao.pptx.",
        12, False, "1C2430", space_after=10,
    )

    callout(
        doc, "Cách dùng.",
        "Mỗi mục là một slide. Đọc mục Lời thoại, không đọc chữ trên slide. "
        "Mục Chỉ vào là chỗ tay hoặc con trỏ nên dừng. Ghi chú của PowerPoint cũng chứa lời thoại này.",
        "E7F7F3",
    )
    callout(
        doc, "Trước buổi nói.",
        "Mở sẵn terminal trong thư mục project và chạy python run.py web. "
        "Mở trình duyệt http://127.0.0.1:8766. Nếu cổng bận, tắt cửa sổ cũ rồi chạy lại. "
        "Phương án dự phòng khi không có trình duyệt: python run.py demo. "
        "Slide không in họ tên. Nếu cần, thêm tên vào câu chào đầu.",
        "FFF6E4", "8A5A00",
    )

    total = sum(s["seconds"] for s in sections)
    short = sum(s["seconds"] for s in sections if s["cut"] == "ngắn")
    dpara(doc, "Hai nhịp thời gian", 16, True, "0B6E62", space_before=8, space_after=4)
    dpara(
        doc,
        f"Nói hết 15 slide: khoảng {total // 60} phút {total % 60} giây, chưa kể demo. "
        f"Nhịp 12 phút chỉ nói các slide đánh dấu Có trong bảng: khoảng {short // 60} phút {short % 60} giây, "
        "rồi demo nhịp 2 khoảng 1 phút. Slide Nếu còn giờ bỏ qua, không tóm tắt từng slide bỏ.",
        12, space_after=8,
    )
    add_timing_table(doc, sections)

    dpara(doc, "Lời thoại theo slide", 16, True, "0B6E62", space_before=12, space_after=6)

    for sec in sections:
        table = doc.add_table(rows=1, cols=1)
        no_table_borders(table)
        cell = table.cell(0, 0)
        shade_cell(cell, "1C2430")
        set_cell_margins(cell, 80, 80, 140, 140)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = DPt(0)
        drun(p, f"SLIDE {sec['n']:02d}   ·   {sec['seconds']} GIÂY   ·   {sec['kicker'].upper()}", 11, True, "8EE0D2")
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_before = DPt(2)
        p2.paragraph_format.space_after = DPt(0)
        drun(p2, sec["title"], 16, True, "FFFFFF")
        doc.add_paragraph().paragraph_format.space_after = DPt(2)
        dpara(doc, "Mục tiêu  ·  " + sec["goal"], 12, False, "5C6B7A", space_after=6)
        callout(doc, "Lời thoại.", sec["say"], "E7F7F3")
        if sec.get("point"):
            dpara(doc, "Chỉ vào  ·  " + sec["point"], 12, False, "1C2430", space_after=4)
        if sec.get("bridge"):
            dpara(doc, "Câu chuyển  ·  " + sec["bridge"], 12, italic=True, color="5C6B7A", space_after=4)
        if sec.get("if_asked"):
            callout(doc, "Nếu bị hỏi ngay.", sec["if_asked"], "FFF6E4", "8A5A00")
        dpara(doc, "", 6, space_after=2)

    dpara(doc, "Demo trực tiếp, khoảng 2–3 phút", 16, True, "0B6E62", space_before=8, space_after=6)
    dpara(
        doc,
        "Không chứng minh lại việc cài đặt. Trang đã mở. Mỗi nhịp: bấm hoặc dán câu, đợi kết quả, nói một câu, rồi chuyển.",
        12, space_after=6,
    )
    demo_steps = [
        ("Nhịp 1 · câu đúng · khoảng 40 giây",
         "Chế độ Hỏi, kho trích Thông tư 08. Bấm câu mẫu nợ tín chỉ.",
         "Kỳ vọng: GGF 100%, nhãn nguồn kéo theo, trích Điều 11 khoản 1 điểm a, trong câu có vượt quá 24.",
         "Nói: khoản này kéo theo câu. GGF là tỉ lệ mệnh đề như vậy. Câu này một mệnh đề, mệnh đề khớp khoản, nên GGF bằng 1."),
        ("Nhịp 2 · nói ngược · khoảng 50 giây",
         "Chuyển chế độ Kiểm một câu nói. Bấm câu mẫu vượt quá 24 thì buộc thôi học.",
         "Kỳ vọng: nhãn nói ngược nguồn. Có dòng cho biết nếu gộp cả điều thì nhãn thành nguồn kéo theo.",
         "Nói: 24 tín chỉ là mức cảnh báo, nằm ở khoản 1. Buộc thôi học nằm ở khoản khác. Chấm từng khoản thì thấy lệch. Gộp cả điều thì hai ý đứng cạnh nhau, dễ gán nhầm là đúng."),
        ("Nhịp 3 · sửa số · khoảng 40 giây",
         "Vẫn ở Kiểm một câu nói. Dán: Thời gian tối đa để sinh viên hoàn thành khóa học không vượt quá 03 lần thời gian theo kế hoạch học tập chuẩn.",
         "Kỳ vọng: nói ngược nguồn, vì khoản ghi 02 lần. Sau một vòng sửa, câu về 02 lần và GGF về 1.",
         "Nói: hệ cố ý đổi 02 thành 03 để đo thước. Verifier bắt lệch. Vòng sửa thay lại đúng khoản. Đây là verify-then-revise, không phải mô hình viết lại."),
    ]
    for title, do, expect, say in demo_steps:
        dpara(doc, title, 13, True, "1C2430", space_before=6, space_after=3)
        dpara(doc, "Làm  ·  " + do, 12, space_after=2)
        dpara(doc, "Kỳ vọng  ·  " + expect, 12, space_after=2)
        callout(doc, "Nói.", say, "FFECEA", "E24B3C")

    dpara(
        doc,
        "Nếu còn hơn một phút: chế độ Hỏi, câu nghỉ học tạm thời. Graph lấy cả Điều 15 và Điều 2 khoản 5, có cụm không vượt quá 02 lần. Naive thường dừng ở một khoản. Nói: câu hai bước là chỗ cạnh dẫn chiếu có việc để làm, khớp với Recall viện dẫn trên slide truy hồi.",
        12, space_after=8,
    )
    callout(
        doc, "Khi trình duyệt không mở.",
        "Chạy python run.py demo. Đọc theo đúng ba nhịp trên. Không gõ thêm. Các số GGF 1,00 và nhãn contradicted nằm trong đầu ra terminal.",
        "FFF6E4", "8A5A00",
    )

    dpara(doc, "Câu hỏi có thể gặp", 16, True, "0B6E62", space_before=10, space_after=6)
    qa = [
        ("Vì sao không dùng mô hình sinh câu trả lời?",
         "Máy buổi này không có GPU. Câu trả lời cố ý là đoạn trích, để đo hai việc tách bạch: có chọn đúng khoản không, và rule có bắt số bị đổi không. Viết lại bằng mô hình là việc còn lại, khi đó mới đo F1 của câu sau sửa."),
        ("HippoRAG 0,633 có phải hơn bài HippoRAG 2 không?",
         "Không so được. Cổng trong bài lan một lớp cạnh trên cùng chỉ mục BM25, không chạy OpenIE và không chạy mã gốc. 0,633 là Recall pool trên 150 câu ViEduReg, không phải F1 của bài được dẫn."),
        ("Macro-F1 bằng 1 có đáng tin không?",
         "Chỉ trên 14 mệnh đề người, đúng các lỗi guideline đã viết rule để bắt. Bỏ rule số thì còn 0,436. Trên mệnh đề dựng từ toàn văn, độ đúng lớp contradicted khoảng 0,15. Em coi 1,000 là kiểm tra rule, không phải kết quả trên câu hỏi tự do."),
        ("Vì sao PhoBERT thấp hơn TF-IDF?",
         "PhoBERT-base tinh chỉnh một epoch, tối đa 800 mẫu tổng hợp. Dev tổng hợp đạt 0,925, nhưng 14 mệnh đề người chỉ 0,419, F1 contradicted 0,222. Mệnh đề tổng hợp là đoạn cắt hoặc đổi số cục bộ. Mệnh đề người còn có lỗi trộn hệ quả. Vài trăm cặp chưa thay được rule."),
        ("GGF tăng sau sửa có nghĩa là câu được viết lại hay hơn không?",
         "Không. Vòng sửa thay mảnh lệch bằng đoạn khoản đang đối chiếu, và bỏ mảnh unsupported. Đoạn thay vào trùng nguồn nên rule gán supported. Em chưa có F1 của câu do mô hình viết."),
        ("RAGAS 0,75 và giám khảo 0,76 là mô hình lớn chấm à?",
         "Không. RAGAS-style trong bài là tỉ lệ mệnh đề được chế độ gộp điều gán supported. Giám khảo trong lần chạy này trùng GGF của rule, hệ số 0,763. Em không gọi RAGAS gốc và không gọi mô hình lớn. Thước đếm từ thuần gần như không tương quan, khoảng âm 0,05."),
        ("Đã có Cohen kappa chưa?",
         "Chưa. 150 câu chờ giảng viên duyệt kép. 127 mệnh đề là một lượt đối chiếu khoản, 72 supported, 26 contradicted, 29 unsupported. 14 mệnh đề test cũng chưa có người thứ hai. Chưa có kappa thì chưa kết luận độ khớp với người nói chung."),
        ("ViHERMES nằm ở đâu trong kết quả?",
         "Demo có kho ViHERMES. Bảng Recall và macro-F1 trong buổi nói lấy trên quy chế đào tạo. Bộ ViHERMES trên Hugging Face cần tài khoản, nên chưa đưa vào bảng. Số F1 trong các bài gốc là mốc tài liệu, không phải baseline của lần chạy này."),
        ("Thông tư 56/2026 đã đối chiếu Công báo chưa?",
         "Toàn văn lấy từ trang pháp luật công khai. Thuộc tính trên cổng Chính phủ xác nhận số hiệu và ngày hiệu lực 7 tháng 7 năm 2026. Đây không phải file Công báo gốc. Thông tư 08/2021 lấy từ PDF Công báo. Thông tư 23/2021 đã hết hiệu lực từ 30 tháng 6 năm 2026."),
        ("Vì sao Graph đôi khi thấp hơn Naive?",
         "Graph đi cạnh từ một hạt giống. Hạt giống sai thì cạnh đi ra cũng sai. Trên 63 câu toàn văn, Recall@5 của graph là 0,333, naive là 0,421. HippoRAG-style lấy ba hạt giống nên viện dẫn lên 0,625, so với 0,458 của naive."),
    ]
    for i, (q, a) in enumerate(qa, 1):
        dpara(doc, f"{i}. {q}", 13, True, "1C2430", space_before=6, space_after=3)
        dpara(doc, a, 12, space_after=4)

    dpara(doc, "Thẻ số, một trang", 16, True, "0B6E62", space_before=12, space_after=6)
    dpara(doc, "Chỉ đọc khi bị hỏi một ô. Không đọc cả thẻ trong phần nói.", 12, False, "5C6B7A", space_after=6)
    facts = [
        ("Đồ thị toàn văn", "545 khoản/điểm, 125 REFERS_TO, 2 REPLACES. 08/2021, 23/2021, 56/2026."),
        ("ViEduReg", "213 khoản/điểm trên 56/2026. Pool sinh 305 câu. Còn 150 câu đã đối chiếu: 90 một khoản, 60 nhiều khoản."),
        ("Recall 150 câu", "HippoRAG pool 0,633, câu trả lời 0,270. Naive pool 0,440, câu trả lời 0,200. Mọi hệ, câu trả lời khoảng 0,20–0,27."),
        ("Câu hai bước", "Pool 0,783–0,983. Câu trả lời khoảng 0,45–0,60. LightRAG-style cao nhất phần câu trả lời hop-2: 0,600."),
        ("Recall@5, 63 câu", "HippoRAG 0,563. Naive 0,421. Graph 0,333. LightRAG-style 0,317. KAG-style 0,405."),
        ("Viện dẫn", "HippoRAG 0,625 so với Naive 0,458. Trích khoản: HippoRAG 0,500, Naive 0,361, Graph 0,250."),
        ("14 mệnh đề", "Triple macro-F1 1,000. Chunk 0,941. Bỏ số 0,436. TF-IDF 0,578. PhoBERT 0,419. AUROC triple trên tập này 1,000."),
        ("PhoBERT", "Dev tổng hợp macro-F1 0,925. Test người 0,419. F1 contradicted trên người 0,222."),
        ("Toàn văn, claim dựng", "Triple: supported 1,00, contradicted 0,15, unsupported 1,00, trung bình 0,717."),
        ("Tương quan", "GGF rule Spearman 1,000. Phủ từ −0,052. RAGAS-style 0,750. Giám khảo = GGF rule 0,763. Không phải RAGAS gốc, không phải mô hình lớn."),
        ("Vòng sửa", "13 câu: GGF 0,231 → 1, contradicted 0,744 → 0. 20 câu toàn văn: GGF 0,45 → 1, contradicted 0,36 → 0."),
        ("127 mệnh đề", "72 supported, 26 contradicted, 29 unsupported. Một lượt. Chưa kappa. Không dùng làm macro-F1 headline."),
        ("Tách huấn luyện", "Train 2.688, dev 298, test 14. Test trùng 14 mệnh đề người. Tập người không trộn vào train."),
        ("Demo", "python run.py web rồi mở http://127.0.0.1:8766. Dự phòng: python run.py demo."),
    ]
    table = doc.add_table(rows=1, cols=2)
    for i, text in enumerate(["Chỗ", "Số cần nhớ"]):
        cell = table.rows[0].cells[i]
        shade_cell(cell, "0E9A86")
        set_cell_margins(cell, 50, 50, 80, 80)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = DPt(0)
        drun(p, text, 11, True, "FFFFFF")
    for i, (k, v) in enumerate(facts):
        row = table.add_row().cells
        fill = "FFF8F3" if i % 2 == 0 else "FFFFFF"
        for cell, text, bold in ((row[0], k, True), (row[1], v, False)):
            shade_cell(cell, fill)
            set_cell_margins(cell, 40, 40, 80, 80)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = DPt(0)
            drun(p, text, 10, bold, "1C2430")

    dpara(
        doc,
        "Nguồn số: BAI_BAO.md, eval/summary.txt, eval/viedureg_systems.txt, eval/public_report.txt. Ngày đối chiếu trong bài: 30.10.2026.",
        10, False, "5C6B7A", space_before=10,
    )
    doc.save(DOCX_PATH)


def _emu_in(value):
    return value / 914400


def audit(pptx_path):
    prs = Presentation(pptx_path)
    warnings = []
    for si, slide in enumerate(prs.slides, 1):
        text_boxes = []
        for sh in slide.shapes:
            x, y = _emu_in(sh.left), _emu_in(sh.top)
            w, h = _emu_in(sh.width), _emu_in(sh.height)
            if x < -0.05 or y < -0.05 or x + w > W + 0.05 or y + h > H + 0.05:
                warnings.append(f"slide {si:02d} shape off-slide ({x:.2f},{y:.2f},{w:.2f},{h:.2f})")
            if not getattr(sh, "has_text_frame", False):
                continue
            paragraphs = []
            for p in sh.text_frame.paragraphs:
                runs = [(r.text or "", r.font.size.pt if r.font.size else 14) for r in p.runs]
                if not runs:
                    continue
                text = "".join(t for t, _ in runs)
                size = max(sz for _, sz in runs)
                if text.strip():
                    paragraphs.append((text, size, p.space_after.pt if p.space_after else 0))
            if not paragraphs:
                continue
            tf = sh.text_frame
            ml = _emu_in(tf.margin_left or 0)
            mr = _emu_in(tf.margin_right or 0)
            mt = _emu_in(tf.margin_top or 0)
            mb = _emu_in(tf.margin_bottom or 0)
            usable_w = max(0.2, w - ml - mr)
            usable_h = max(0.1, h - mt - mb)
            need = 0
            for text, size, after in paragraphs:
                cpl = max(1, usable_w / (size * 0.47 / 72))
                lines = 0
                for line in text.split("\n"):
                    lines += max(1, int((len(line) + cpl - 1) // cpl))
                need += lines * (size * 1.12 / 72) + after / 72
            if need > usable_h + 0.08:
                sample = paragraphs[0][0][:42].replace("\n", " ")
                warnings.append(
                    f"slide {si:02d} overflow ~{need:.2f}in in {usable_h:.2f}in box: {sample}"
                )
            text_boxes.append((x, y, w, h, paragraphs[0][0][:28].replace("\n", " ")))
        for i, a in enumerate(text_boxes):
            for b in text_boxes[i + 1 :]:
                ox = min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0])
                oy = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
                if ox > 0.08 and oy > 0.08:
                    warnings.append(
                        f"slide {si:02d} text overlap {ox:.2f}x{oy:.2f}: [{a[4]}] <> [{b[4]}]"
                    )
    return warnings


def main():
    sections = build_pptx()
    assert len(sections) == TOTAL
    build_docx(sections)
    warnings = audit(PPTX_PATH)
    if warnings:
        print(f"AUDIT {len(warnings)}")
        for line in warnings:
            print(line)
    else:
        print("AUDIT clean")
    talk = sum(s["seconds"] for s in sections)
    short = sum(s["seconds"] for s in sections if s["cut"] == "ngắn")
    print(f"pptx {PPTX_PATH}")
    print(f"docx {DOCX_PATH}")
    print(f"slides {len(sections)}  talk {talk}s  short {short}s")


if __name__ == "__main__":
    main()
