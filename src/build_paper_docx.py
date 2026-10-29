"""Xuất BAI_BAO.md thành Word, Times New Roman, giãn dòng 1,5."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "BAI_BAO.md"
DST = ROOT / "BAI_BAO_ViKAG_Faith.docx"


def _font(run, size, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def _body(p, first_line=True):
    fmt = p.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(3)
    fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if first_line:
        fmt.first_line_indent = Cm(1.0)


def add_runs(p, text, size, bold=False, italic=False):
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            _font(run, size, bold=True, italic=italic)
        else:
            run = p.add_run(part)
            _font(run, size, bold=bold, italic=italic)


def add_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.autofit = True
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            cell_p = table.rows[i].cells[j].paragraphs[0]
            cell_p.paragraph_format.space_before = Pt(0)
            cell_p.paragraph_format.space_after = Pt(0)
            cell_p.paragraph_format.line_spacing = 1.0
            run = cell_p.add_run(cell)
            _font(run, 11, bold=(i == 0))
    doc.add_paragraph()


def parse_table(block: str) -> list[list[str]]:
    rows = []
    for line in block.splitlines():
        if not line.strip().startswith("|"):
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if all(set(c) <= set("-: ") for c in cells):
            continue
        rows.append(cells)
    return rows


def main() -> None:
    text = SRC.read_text(encoding="utf-8")
    lines = text.splitlines()
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(13)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    i = 0
    title_done = False
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            add_runs(p, line[2:].strip(), 16, bold=True)
            title_done = True
            i += 1
            continue
        if line.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            add_runs(p, line[3:].strip(), 14, bold=True)
            i += 1
            continue
        if line.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            add_runs(p, line[4:].strip(), 13, bold=True)
            i += 1
            continue
        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, parse_table("\n".join(block)))
            continue
        if line.startswith("**Bài báo") or line.startswith("Họ và tên") or line.startswith("Môn:") or line.startswith("Ngày:"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            add_runs(p, line.strip(), 13, italic=line.startswith("**Bài"))
            i += 1
            continue
        para = [line.strip()]
        i += 1
        while i < len(lines) and lines[i].strip() and not lines[i].startswith("#") and not lines[i].startswith("|"):
            para.append(lines[i].strip())
            i += 1
        body = " ".join(para)
        p = doc.add_paragraph()
        abstract = body.startswith("Hỏi đáp trên văn bản") or body.startswith("**Từ khóa:**")
        _body(p, first_line=not abstract and title_done)
        if abstract and body.startswith("Hỏi đáp"):
            p.paragraph_format.first_line_indent = Cm(0)
        add_runs(p, body, 13, italic=body.startswith("Hỏi đáp trên văn bản"))
    doc.save(DST)
    print(DST)


if __name__ == "__main__":
    main()
